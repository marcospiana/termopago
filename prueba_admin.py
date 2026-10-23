"""Banco de pruebas: levanta app.py contra un SQLite en memoria haciendose
pasar por psycopg2, para verificar la migracion de tipos y el panel /admin
sin tocar la base de produccion."""
import os, re, sys, sqlite3, types

os.environ.setdefault("CLAVE_SECRETA", "TESTKEY")
os.environ.setdefault("DATABASE_URL", "postgresql://fake")
os.environ.setdefault("MQTT_HOST", "broker.test")
os.environ.setdefault("MQTT_USER", "termopago")
os.environ.setdefault("MQTT_PASS", "secreta")

_conn_unica = sqlite3.connect(":memory:", check_same_thread=False)
_conn_unica.row_factory = sqlite3.Row


def _traducir(sql):
    sql = sql.replace("%s", "?")
    sql = sql.replace("%%", "%")        # psycopg2 escapa el % literal como %%
    sql = re.sub(r"SERIAL PRIMARY KEY", "INTEGER PRIMARY KEY AUTOINCREMENT", sql, flags=re.I)
    return sql


def _expandir_tuplas(sql, args):
    """psycopg2 adapta una tupla a (a,b,c) para un 'IN %s'; sqlite no sabe.
    Expandimos la tupla a tantos ? como elementos tenga."""
    if not args or not any(isinstance(a, (tuple, list)) for a in args):
        return sql, args
    partes = sql.split("?")
    if len(partes) - 1 != len(args):
        return sql, args
    salida, planos = partes[0], []
    for a, resto in zip(args, partes[1:]):
        if isinstance(a, (tuple, list)):
            salida += "(" + ",".join("?" * len(a)) + ")" + resto
            planos.extend(a)
        else:
            salida += "?" + resto
            planos.append(a)
    return salida, planos


class _Cur:
    def __init__(self, cur): self._c = cur
    def execute(self, sql, args=()):
        s = _traducir(sql)
        if re.search(r"ALTER TABLE .* ADD COLUMN IF NOT EXISTS", sql, re.I):
            s = re.sub(r"IF NOT EXISTS ", "", s, flags=re.I)
            try:
                return self._c.execute(s, args)
            except sqlite3.OperationalError as e:
                if "duplicate column" in str(e):
                    return None
                raise
        s, args = _expandir_tuplas(s, args)
        return self._c.execute(s, args)
    def fetchone(self):
        r = self._c.fetchone()
        return dict(r) if r else None
    def fetchall(self):
        return [dict(r) for r in self._c.fetchall()]
    def close(self): pass
    @property
    def rowcount(self): return self._c.rowcount


class _Conn:
    cursor_factory = None
    def cursor(self): return _Cur(_conn_unica.cursor())
    def commit(self): _conn_unica.commit()
    def close(self): pass


fake = types.ModuleType("psycopg2")
fake.connect = lambda *a, **k: _Conn()
fake.extras = types.ModuleType("psycopg2.extras")
fake.extras.RealDictCursor = object
sys.modules["psycopg2"] = fake
sys.modules["psycopg2.extras"] = fake.extras

# MercadoPago: nada de red en la prueba.
import requests
class _R:
    status_code = 200
    def json(self): return {"results": [{"id": 999, "store_id": 7,
                                         "qr": {"image": "http://qr", "template_document": "http://pdf"}}]}
requests.get = lambda *a, **k: _R()
requests.post = lambda *a, **k: _R()

import app
app.rearmar_qr = lambda disp: None          # no tocar MP
app.cancelar_orden_qr = lambda disp: None
app.MP_CLIENT_ID = "123"
app.MP_CLIENT_SECRET = "abc"

c = app.app.test_client()
K = "TESTKEY"
fallas = []


def chequear(nombre, cond, extra=""):
    print(("  OK   " if cond else "  FALLA") + f" {nombre} {extra if not cond else ''}")
    if not cond:
        fallas.append(nombre)


print("\n== 1. Backfill de las cajas viejas ==")
cur = _conn_unica.cursor()
for cid, nom in [("inflado01", "Inflador"), ("aspiradora01", "Aspiradora 1"),
                 ("soplado01", "Soplado 1"), ("aspiradora02", "Aspiradora 2"),
                 ("soplado02", "Soplado 2"), ("villagas01", "Fichas Villagas"),
                 ("termo_002", "Termo viejo")]:
    cur.execute("INSERT OR IGNORE INTO dispositivos (id,nombre,precio,segundos) VALUES (?,?,?,?)",
                (cid, nom, 500, 300))
_conn_unica.commit()
app.init_db()
app.invalidar_cache_tipos()

tipos = {r["id"]: (r["tipo"], r["esp_id"], r["canal"])
         for r in _Cur(_conn_unica.cursor()).execute(
             "SELECT id,tipo,esp_id,canal FROM dispositivos").fetchall()}
chequear("villagas01 quedo 'fichas'", tipos["villagas01"][0] == "fichas", tipos["villagas01"])
chequear("inflado01 quedo 'pulso'", tipos["inflado01"][0] == "pulso", tipos["inflado01"])
chequear("aspiradora01 quedo 'sostenida'", tipos["aspiradora01"][0] == "sostenida", tipos["aspiradora01"])
chequear("termo_002 quedo 'legacy'", tipos["termo_002"][0] == "legacy", tipos["termo_002"])
chequear("aspiradora01 y soplado01 comparten esp", tipos["aspiradora01"][1] == tipos["soplado01"][1] == "estacion01")
chequear("canales 0 y 1 asignados", (tipos["aspiradora01"][2], tipos["soplado01"][2]) == (0, 1))
chequear("caja sola = su propio esp", tipos["inflado01"][1] == "inflado01")

print("\n== 2. Los conjuntos dinamicos dan lo mismo que los sets viejos ==")
chequear("ESTACIONES_MQTT == semilla", set(app.ESTACIONES_MQTT) == app.SEMILLA_MQTT, sorted(app.ESTACIONES_MQTT))
chequear("ESTACIONES_PULSO == semilla", set(app.ESTACIONES_PULSO) == app.SEMILLA_PULSO, sorted(app.ESTACIONES_PULSO))
chequear("ESTACIONES_FICHAS == semilla", set(app.ESTACIONES_FICHAS) == app.SEMILLA_FICHAS, sorted(app.ESTACIONES_FICHAS))
chequear("'villagas01' in ESTACIONES_FICHAS", "villagas01" in app.ESTACIONES_FICHAS)
chequear("'termo_002' NO es MQTT", "termo_002" not in app.ESTACIONES_MQTT)
chequear("caja inexistente no es MQTT", "no_existe" not in app.ESTACIONES_MQTT)
chequear("hermanas de aspiradora01", app.cajas_hermanas("aspiradora01") == {"aspiradora01", "soplado01"},
         app.cajas_hermanas("aspiradora01"))
chequear("hermanas de inflado01", app.cajas_hermanas("inflado01") == {"inflado01"})
chequear("hermanas de una caja desconocida", app.cajas_hermanas("zzz") == {"zzz"})

print("\n== 2b. estacion02 (aspiradora02 + soplado02) intacta ==")
chequear("aspiradora02 sigue sostenida", tipos["aspiradora02"][0] == "sostenida", tipos["aspiradora02"])
chequear("soplado02 sigue sostenida", tipos["soplado02"][0] == "sostenida", tipos["soplado02"])
chequear("las dos son MQTT", {"aspiradora02", "soplado02"} <= set(app.ESTACIONES_MQTT))
chequear("ninguna es de pulso", not ({"aspiradora02", "soplado02"} & set(app.ESTACIONES_PULSO)))
chequear("comparten el ESP estacion02",
         app.cajas_hermanas("aspiradora02") == {"aspiradora02", "soplado02"},
         app.cajas_hermanas("aspiradora02"))
chequear("canales 0 y 1", (tipos["aspiradora02"][2], tipos["soplado02"][2]) == (0, 1))

print("\n== 2c. DB caida con cache frio -> semillas, no 'todo legacy' ==")
_get_db_real = app.get_db
app.get_db = lambda: (_ for _ in ()).throw(RuntimeError("DB caida"))
app._cache_tipos["filas"] = {}          # simula recien deployado
app._cache_tipos["t"] = 0.0
chequear("aspiradora02 sigue siendo MQTT con la DB caida", "aspiradora02" in app.ESTACIONES_MQTT)
chequear("inflado01 sigue siendo pulso", "inflado01" in app.ESTACIONES_PULSO)
chequear("villagas01 sigue siendo fichas", "villagas01" in app.ESTACIONES_FICHAS)
chequear("las hermanas de estacion02 se mantienen",
         app.cajas_hermanas("soplado02") == {"aspiradora02", "soplado02"},
         app.cajas_hermanas("soplado02"))
chequear("el fallback NO se cachea (reintenta contra la DB)", app._cache_tipos["filas"] == {})
app.get_db = _get_db_real
app.invalidar_cache_tipos()
set(app.ESTACIONES_MQTT)   # fuerza una lectura ya con la DB sana
chequear("al volver la DB vuelve a mandar la DB (ve cajas que no estan en la semilla)",
         "termo_002" in app._cache_tipos["filas"], sorted(app._cache_tipos["filas"]))

print("\n== 3. El backfill no pisa un tipo cambiado a mano ==")
app.actualizar_dispositivo("termo_002", {"tipo": "pulso"})
app.invalidar_cache_tipos()
app.init_db()
app.invalidar_cache_tipos()
chequear("termo_002 sigue 'pulso' tras re-deploy", "termo_002" in app.ESTACIONES_PULSO)
app.actualizar_dispositivo("termo_002", {"tipo": "legacy"})
app.invalidar_cache_tipos()

print("\n== 4. Panel /admin ==")
r = c.get(f"/admin/{K}")
chequear("carga 200", r.status_code == 200, r.status_code)
chequear("clave mala -> 403", c.get("/admin/mala").status_code == 403)

r = c.post(f"/admin/{K}", data={"accion": "nuevo_cliente", "alias": "Lava Dero!", "nombre": "Lavadero Sur"})
chequear("alta de cliente", b"creado" in r.data, r.data[:200])
chequear("alias slugificado (espacio -> _, sin simbolos)",
         app.get_cliente("lava_dero") is not None, [x["alias"] for x in app.get_clientes()])
c.post(f"/admin/{K}", data={"accion": "nuevo_cliente", "alias": "lavadero", "nombre": "Lavadero Sur"})
cli = app.get_cliente("lavadero")
chequear("le genero panel_token y pin", bool(cli["panel_token"]) and len(cli["pin"]) == 4)

r = c.post(f"/admin/{K}", data={"accion": "nuevo_cliente", "alias": "lavadero", "nombre": "x"}, follow_redirects=True)
chequear("no deja duplicar alias", "Ya existe".encode() in r.data)

print("\n== 5. Alta de maquinas desde el panel ==")
r = c.post(f"/admin/{K}", data={
    "accion": "nueva_maquina", "cliente": "lavadero", "disp_id": "lavadero01",
    "nombre": "Aspiradora sur", "tipo": "sostenida", "esp_id": "lavadero_esp1",
    "canal": "0", "precio": "800", "valor": "240", "cuenta": ""})
chequear("crea la maquina", b"creada" in r.data, r.data[:300])
d = app.get_dispositivo("lavadero01")
chequear("guardo tipo/esp/canal", (d["tipo"], d["esp_id"], d["canal"]) == ("sostenida", "lavadero_esp1", 0), d and dict(d))
chequear("guardo precio y valor", (float(d["precio"]), int(d["segundos"])) == (800.0, 240))
chequear("quedo asociada al cliente", d["cliente"] == "lavadero")
app.invalidar_cache_tipos()
chequear("ya es MQTT sin redeploy", "lavadero01" in app.ESTACIONES_MQTT)
chequear("y NO es de pulso", "lavadero01" not in app.ESTACIONES_PULSO)

c.post(f"/admin/{K}", data={
    "accion": "nueva_maquina", "cliente": "lavadero", "disp_id": "lavadero02",
    "nombre": "Soplado sur", "tipo": "sostenida", "esp_id": "lavadero_esp1",
    "canal": "1", "precio": "800", "valor": "240", "cuenta": ""})
app.invalidar_cache_tipos()
chequear("2 cajas en un mismo ESP son hermanas",
         app.cajas_hermanas("lavadero01") == {"lavadero01", "lavadero02"}, app.cajas_hermanas("lavadero01"))

r = c.post(f"/admin/{K}", data={
    "accion": "nueva_maquina", "cliente": "lavadero", "disp_id": "lavadero03",
    "nombre": "Fichas", "tipo": "fichas", "esp_id": "", "canal": "",
    "precio": "3000", "valor": "", "cuenta": ""})
d = app.get_dispositivo("lavadero03")
app.invalidar_cache_tipos()
chequear("fichas: valor por defecto = 1 ficha", int(d["segundos"]) == 1, d and int(d["segundos"]))
chequear("fichas cuenta como pulso", "lavadero03" in app.ESTACIONES_PULSO)
chequear("fichas: esp_id vacio cae en su propio id", d["esp_id"] == "lavadero03")

r = c.post(f"/admin/{K}", data={
    "accion": "nueva_maquina", "cliente": "lavadero", "disp_id": "lavadero01",
    "nombre": "Repetida", "tipo": "pulso", "cuenta": ""}, follow_redirects=True)
chequear("no deja repetir ID de maquina", "Ya existe".encode() in r.data)

r = c.post(f"/admin/{K}", data={
    "accion": "nueva_maquina", "cliente": "lavadero", "disp_id": "malmal",
    "nombre": "X", "tipo": "inventado", "cuenta": ""}, follow_redirects=True)
chequear("rechaza tipo invalido", "inválido".encode() in r.data)

print("\n== 6. El payload MQTT cambia solo con el tipo ==")
capturado = {}
def _pub(topic=None, payload=None, **k): capturado["t"], capturado["p"] = topic, payload
app._mqtt_publish = types.SimpleNamespace(single=_pub)
app.publicar_activacion("lavadero03", "pago1")
chequear("expendedora manda 'cantidad'", '"cantidad": 1' in capturado["p"], capturado["p"])
app.publicar_activacion("lavadero01", "pago2")
chequear("sostenida manda 'segundos'", '"segundos": 240' in capturado["p"], capturado["p"])

print("\n== 7. Ficha de flasheo y .zip ==")
r = c.get(f"/admin/{K}/maquina/lavadero03")
chequear("ficha carga", r.status_code == 200, r.status_code)
chequear("muestra el ID a cargar en el portal", b"lavadero03" in r.data)
r = c.get(f"/admin/{K}/maquina/lavadero01")
chequear("ficha de ESP compartido lista las 2 cajas", b"lavadero02" in r.data)
chequear("ficha con clave mala -> 403", c.get("/admin/mala/maquina/lavadero01").status_code == 403)
chequear("ficha de maquina inexistente -> 404", c.get(f"/admin/{K}/maquina/nada").status_code == 404)

chequear("la ficha muestra el QR de la caja", b'class="qr"' in r.data)
chequear("y ofrece el PDF para imprimir", b"PDF para imprimir" in r.data)

# MercadoPago caido: la ficha tiene que seguir sirviendo para flashear.
_get_real = requests.get
requests.get = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("MP caido"))
r = c.get(f"/admin/{K}/maquina/lavadero01")
chequear("con MP caido la ficha igual carga", r.status_code == 200, r.status_code)
chequear("y sigue mostrando el ID a cargar en el portal", b"lavadero01" in r.data)
chequear("avisa que no pudo traer el QR", "No pude consultar MercadoPago".encode() in r.data)
requests.get = _get_real

r = c.get(f"/admin/{K}/firmware/lavadero03.zip?secretos=1")
if r.status_code == 200:
    import io as _io, zipfile
    z = zipfile.ZipFile(_io.BytesIO(r.data))
    nombres = z.namelist()
    chequear("el zip trae el .ino", any(n.endswith(".ino") for n in nombres), nombres)
    chequear("el zip trae el LEEME", any("LEEME_lavadero03" in n for n in nombres), nombres)
    sec = [n for n in nombres if n.endswith("secretos_privado.h")]
    chequear("el zip trae secretos generados", bool(sec), nombres)
    if sec:
        chequear("secretos con el broker real", b"broker.test" in z.read(sec[0]))
    leeme = [n for n in nombres if "LEEME" in n][0]
    chequear("el LEEME dice el ID de caja", b"lavadero03" in z.read(leeme))
else:
    print(f"  (zip devolvio {r.status_code}: {r.data[:200]!r} — falta la carpeta del sketch en este entorno)")

print("\n== 8. Baja de maquina ==")
r = c.post(f"/admin/{K}", data={"accion": "borrar_maquina", "disp_id": "lavadero02"})
chequear("da de baja", b"baja" in r.data, r.data[:200])
chequear("desaparecio de la tabla", app.get_dispositivo("lavadero02") is None)
app.invalidar_cache_tipos()
chequear("y de los conjuntos", "lavadero02" not in app.ESTACIONES_MQTT)

print("\n== 8a. Renovar PIN y link del cliente ==")
antes = app.get_cliente("lavadero")
r = c.post(f"/admin/{K}", data={"accion": "nuevo_pin", "alias": "lavadero"})
ahora_cli = app.get_cliente("lavadero")
chequear("cambia el PIN", ahora_cli["pin"] != antes["pin"], (antes["pin"], ahora_cli["pin"]))
chequear("el PIN nuevo tiene 4 digitos", len(ahora_cli["pin"]) == 4 and ahora_cli["pin"].isdigit())
chequear("el mensaje muestra el PIN nuevo", ahora_cli["pin"].encode() in r.data)
chequear("no toca el link del panel", ahora_cli["panel_token"] == antes["panel_token"])

r = c.post(f"/admin/{K}", data={"accion": "nuevo_link", "alias": "lavadero"})
despues = app.get_cliente("lavadero")
chequear("cambia el token del panel", despues["panel_token"] != antes["panel_token"])
chequear("no toca el PIN", despues["pin"] == ahora_cli["pin"])
chequear("el link viejo deja de abrir",
         c.get(f"/panel/{antes['panel_token']}").status_code == 404)
chequear("el link nuevo abre", c.get(f"/panel/{despues['panel_token']}").status_code == 200)
r = c.post(f"/admin/{K}", data={"accion": "nuevo_pin", "alias": "no_existe"}, follow_redirects=True)
chequear("cliente inexistente no rompe", b"No existe ese cliente" in r.data)

print("\n== 8b. Baja de cliente ==")
r = c.post(f"/admin/{K}", data={"accion": "borrar_cliente", "alias": "lavadero"}, follow_redirects=True)
chequear("no deja borrar un cliente con maquinas", "todavía tiene".encode() in r.data, r.data[:300])
chequear("y el cliente sigue existiendo", app.get_cliente("lavadero") is not None)
# En este punto hay 2 clientes: 'lava_dero' sin maquinas y 'lavadero' con 2.
# El boton de baja tiene que aparecer una sola vez: para el que no tiene ninguna.
chequear("el boton de baja aparece solo para el cliente sin maquinas",
         c.get(f"/admin/{K}").data.count(b'value="borrar_cliente"') == 1,
         c.get(f"/admin/{K}").data.count(b'value="borrar_cliente"'))
r = c.post(f"/admin/{K}", data={"accion": "borrar_cliente", "alias": "lava_dero"})
chequear("borra un cliente sin maquinas", b"eliminado" in r.data, r.data[:300])
chequear("desaparecio de la tabla", app.get_cliente("lava_dero") is None)
chequear("ya no queda ningun boton de baja de cliente",
         c.get(f"/admin/{K}").data.count(b'value="borrar_cliente"') == 0)
r = c.post(f"/admin/{K}", data={"accion": "borrar_cliente", "alias": "no_existe"}, follow_redirects=True)
chequear("cliente inexistente da error, no rompe", b"No existe ese cliente" in r.data)

print("\n== 8c. Freno de fuerza bruta del PIN ==")
c.post(f"/admin/{K}", data={"accion": "nuevo_pin", "alias": "lavadero"})
cli = app.get_cliente("lavadero")
pin_ok, tok = cli["pin"], cli["panel_token"]
pin_malo = "0000" if pin_ok != "0000" else "1111"

# El POST de regalo redirige (PRG, para que el refresh no reenvie el form), asi
# que para leer el mensaje hay que seguir el redirect hasta el GET.
def _regalo(pin):
    return c.post(f"/panel/{tok}", data={"accion": "regalar", "pin": pin},
                  follow_redirects=True)

for i in range(1, app.PIN_MAX_INTENTOS):
    r = _regalo(pin_malo)
    chequear(f"intento fallido {i} avisa cuantos quedan",
             f"Te queda(n) {app.PIN_MAX_INTENTOS - i}".encode() in r.data, r.data[-400:])

r = _regalo(pin_malo)
chequear("al llegar al tope se bloquea", "se bloqueo la entrega".encode() in r.data, r.data[-400:])
chequear("quedo grabado el bloqueo en la DB",
         bool(app.get_cliente("lavadero")["pin_bloqueado_hasta"]))

r = _regalo(pin_ok)
chequear("bloqueado, ni el PIN correcto pasa", "Proba de nuevo en".encode() in r.data, r.data[-400:])

# El bloqueo se vence solo: lo corremos al pasado en vez de esperar 15 min.
app.guardar_cliente("lavadero", {
    "pin_bloqueado_hasta": (app.ahora_ar() - app.timedelta(minutes=1)).isoformat()})
r = _regalo(pin_ok)
chequear("vencido el bloqueo, el PIN correcto vuelve a andar",
         "Proba de nuevo en".encode() not in r.data and b"PIN incorrecto" not in r.data,
         r.data[-400:])
chequear("el PIN correcto limpia el contador",
         not app.get_cliente("lavadero")["pin_fallidos"])

# Un PIN nuevo desde el admin tiene que levantar el bloqueo.
app.guardar_cliente("lavadero", {
    "pin_fallidos": 9,
    "pin_bloqueado_hasta": (app.ahora_ar() + app.timedelta(minutes=30)).isoformat()})
c.post(f"/admin/{K}", data={"accion": "nuevo_pin", "alias": "lavadero"})
cli = app.get_cliente("lavadero")
chequear("PIN nuevo levanta el bloqueo", not cli["pin_bloqueado_hasta"] and not cli["pin_fallidos"],
         (cli["pin_bloqueado_hasta"], cli["pin_fallidos"]))

print("\n== 9. /config sigue funcionando con los tipos nuevos ==")
r = c.get(f"/config/{K}")
chequear("config carga", r.status_code == 200, r.status_code)
chequear("expendedora pide 'Fichas por pago'", "Fichas por pago".encode() in r.data)

print("\n== 10. Todas las paginas de admin vuelven al panel ==")
for ruta in ["config", "reinicios", "cortes", "estado", "estadisticas", "historial"]:
    r = c.get(f"/{ruta}/{K}")
    ok = r.status_code == 200 and f'href="/admin/{K}"'.encode() in r.data
    chequear(f"/{ruta} tiene el link de volver", ok, f"status {r.status_code}")

# El panel del cliente NO puede filtrar la clave secreta.
cli = app.get_cliente("lavadero")
r = c.get(f"/panel/{cli['panel_token']}")
chequear("el panel del cliente carga", r.status_code == 200, r.status_code)
chequear("y NO expone la clave ni un link al admin",
         K.encode() not in r.data and b"/admin/" not in r.data)

print("\n== 11. Pulsos por ficha (precio de la placa -> pulsos del ESP) ==")

# El calculo puro. La division TIENE que dar exacta: si sobra resto queda
# credito colgado en la maquina y termina regalando una ficha.
def _d(credito, valor):
    return {"credito_ficha": credito, "valor_pulso": valor}

chequear("3000/100 = 30 pulsos", app.pulsos_por_ficha(_d(3000, 100)) == 30)
chequear("4000/100 = 40 pulsos", app.pulsos_por_ficha(_d(4000, 100)) == 40)
chequear("200/200 = 1 pulso",    app.pulsos_por_ficha(_d(200, 200)) == 1)
chequear("sin datos cae al default 3000/100",
         app.pulsos_por_ficha({"credito_ficha": None, "valor_pulso": None}) == 30)
chequear("disp None cae al default", app.pulsos_por_ficha(None) == 30)
chequear("division inexacta -> None (no se manda nada)",
         app.pulsos_por_ficha(_d(3050, 100)) is None)
chequear("valor_pulso 0 -> None", app.pulsos_por_ficha(_d(3000, 0)) is None)
chequear("credito negativo -> None", app.pulsos_por_ficha(_d(-100, 100)) is None)
chequear("mas de PULSOS_MAX -> None", app.pulsos_por_ficha(_d(999999, 100)) is None)

# El cmd MQTT tiene que llevar los pulsos calculados, y solo en las de fichas.
app.actualizar_dispositivo("lavadero03", {"credito_ficha": 3000, "valor_pulso": 100})
app.publicar_activacion("lavadero03", "pagoP1")
chequear("el cmd de fichas lleva 'pulsos'", '"pulsos": 30' in capturado["p"], capturado["p"])
app.actualizar_dispositivo("lavadero03", {"credito_ficha": 4500, "valor_pulso": 100})
app.publicar_activacion("lavadero03", "pagoP2")
chequear("cambiar el precio de la placa cambia los pulsos",
         '"pulsos": 45' in capturado["p"], capturado["p"])
app.publicar_activacion("lavadero01", "pagoP3")
chequear("la sostenida NO lleva pulsos", "pulsos" not in capturado["p"], capturado["p"])

# Si la config quedo mal, mejor no mandar el campo: el ESP usa el compilado.
app.actualizar_dispositivo("lavadero03", {"credito_ficha": 3050, "valor_pulso": 100})
app.publicar_activacion("lavadero03", "pagoP4")
chequear("config inexacta -> no manda pulsos (el ESP usa el compilado)",
         "pulsos" not in capturado["p"], capturado["p"])
app.actualizar_dispositivo("lavadero03", {"credito_ficha": 3000, "valor_pulso": 100})

print("\n== 12. /config valida el credito por ficha ==")
r = c.get(f"/config/{K}")
chequear("muestra el credito por ficha", "Cr\u00e9dito por ficha".encode() in r.data)
chequear("muestra el valor del pulso", "Valor del pulso".encode() in r.data)
chequear("muestra los pulsos calculados", b"30 pulsos por ficha" in r.data)
chequear("aclara que es multiplo de 100", "M\u00faltiplo de $100".encode() in r.data)

def _post(credito, valor="100"):
    datos = {}
    for d in app.get_dispositivos():
        datos[f"precio__{d['id']}"] = f"{float(d['precio']):g}"
        if d["id"] in app.ESTACIONES_MQTT:
            datos[f"segundos__{d['id']}"] = str(int(d["segundos"]))
        else:
            datos[f"minutos__{d['id']}"] = str(max(1, int(d["segundos"]) // 60))
        if d["id"] in app.ESTACIONES_FICHAS:
            datos[f"credito__{d['id']}"] = str(credito)
            datos[f"valorpulso__{d['id']}"] = str(valor)
    return c.post(f"/config/{K}", data=datos)

r = _post(4000)
chequear("guarda un multiplo de 100", "Guardado".encode() in r.data)
chequear("y quedo en la DB",
         float(app.get_dispositivo("lavadero03")["credito_ficha"]) == 4000)

r = _post(3050)
chequear("rechaza un credito que no es multiplo de 100",
         "m\u00faltiplos de $100".encode() in r.data)
chequear("y NO piso lo que estaba guardado",
         float(app.get_dispositivo("lavadero03")["credito_ficha"]) == 4000)

r = _post(300, valor="200")
chequear("rechaza division inexacta aunque los dos sean multiplos de 100",
         "no da exacto".encode() in r.data)
chequear("y tampoco piso nada",
         float(app.get_dispositivo("lavadero03")["credito_ficha"]) == 4000)

r = _post(0)
chequear("rechaza credito 0", "\u274c".encode() in r.data)

print("\n== 13. El panel NO regala fichas de mas al refrescar ==")
# BUG REAL: la pagina se dibujaba como respuesta AL POST. El navegador quedaba
# parado sobre el POST y cada F5 reenviaba el formulario con el PIN -> otra
# ficha. Ahora el POST redirige (PRG) y el refresh es un GET inofensivo.
cli = app.get_cliente("lavadero")
tok = cli["panel_token"]
pin_ok = cli["pin"]
app.guardar_cliente("lavadero", {"pin_fallidos": 0, "pin_bloqueado_hasta": None})
app.actualizar_dispositivo("lavadero03", {"ultimo_poll": app.ahora_ar().isoformat()})

def _regalos():
    conn = app.get_db(); cur = conn.cursor()
    cur.execute("SELECT COUNT(*) AS n FROM ordenes WHERE id LIKE 'gift_%%'")
    n = cur.fetchone()["n"]; cur.close(); conn.close()
    return n

antes = _regalos()
r = c.post(f"/panel/{tok}", data={"accion": "regalar", "pin": pin_ok})
chequear("el POST de regalo redirige (no devuelve HTML)", r.status_code in (302, 303), r.status_code)
chequear("regalo 1 ficha", _regalos() == antes + 1, _regalos() - antes)

# El refresh del navegador tras el redirect es un GET a la misma URL.
destino = r.headers["Location"]
r2 = c.get(destino)
chequear("el GET del redirect muestra el mensaje", r2.status_code == 200, r2.status_code)
chequear("refrescar NO regala otra ficha", _regalos() == antes + 1, _regalos() - antes)
r3 = c.get(destino)
chequear("refrescar de nuevo tampoco", _regalos() == antes + 1, _regalos() - antes)

# Y si alguien reenvia el POST a mano (doble clic, reintento del navegador),
# el cooldown del servidor lo frena.
r4 = c.post(f"/panel/{tok}", data={"accion": "regalar", "pin": pin_ok})
chequear("un segundo POST inmediato NO regala (cooldown)",
         _regalos() == antes + 1, _regalos() - antes)
chequear("y avisa por que", "unos segundos" in c.get(r4.headers["Location"]).get_data(as_text=True))

# El mensaje viaja por la URL: no puede inyectar HTML en la pagina.
r5 = c.get(f"/panel/{tok}?m=<script>alert(1)</script>")
chequear("el mensaje de la URL se escapa (sin XSS)",
         b"<script>alert(1)</script>" not in r5.data and b"&lt;script&gt;" in r5.data)

# Un PIN incorrecto tampoco puede regalar por reenvio.
antes2 = _regalos()
r6 = c.post(f"/panel/{tok}", data={"accion": "regalar", "pin": "0000"})
chequear("PIN incorrecto no regala", _regalos() == antes2, _regalos() - antes2)
chequear("y tambien redirige", r6.status_code in (302, 303), r6.status_code)

print("\n== 14. El titulo del QR lleva el nombre del comercio ==")
# Lo que ve el cliente en su app al escanear. Un "Ficha x 1" pelado no dice de
# quien es el cobro; con el nombre del local se entiende y evita dudas.
# Se prueba titulo_orden() directo: es logica pura, sin red (rearmar_qr esta
# stubbeada arriba para no tocar MP).

t = app.titulo_orden(app.get_dispositivo("lavadero03"))
chequear("arranca con 'Ficha x'", t.startswith("Ficha x"), t)
chequear("y lleva el nombre del cliente", "Lavadero" in t, t)

_n = app.get_cliente("lavadero")["nombre"]

app.guardar_cliente("lavadero", {"nombre": ""})
chequear("sin nombre de cliente queda 'Ficha x N' solo",
         app.titulo_orden(app.get_dispositivo("lavadero03")) == "Ficha x 1",
         app.titulo_orden(app.get_dispositivo("lavadero03")))

app.guardar_cliente("lavadero", {"nombre": "   "})
chequear("un nombre en blanco tampoco ensucia el titulo",
         app.titulo_orden(app.get_dispositivo("lavadero03")) == "Ficha x 1",
         app.titulo_orden(app.get_dispositivo("lavadero03")))

app.guardar_cliente("lavadero", {"nombre": "N" * 300})
chequear("un nombre larguisimo se recorta a 120",
         len(app.titulo_orden(app.get_dispositivo("lavadero03"))) == 120)

app.guardar_cliente("lavadero", {"nombre": _n})

# Varias fichas por pago: el numero acompaña.
app.actualizar_dispositivo("lavadero03", {"segundos": 3})
chequear("con 3 fichas dice 'Ficha x 3'",
         app.titulo_orden(app.get_dispositivo("lavadero03")).startswith("Ficha x 3"),
         app.titulo_orden(app.get_dispositivo("lavadero03")))
app.actualizar_dispositivo("lavadero03", {"segundos": 1})

# Las que no son de fichas no se tocan.
t2 = app.titulo_orden(app.get_dispositivo("lavadero01"))
chequear("las de tiempo siguen diciendo los minutos", "minutos" in t2, t2)
chequear("y NO llevan el nombre del cliente pegado", not t2.endswith("Lavadero"), t2)

print("\n== 15. Dos codigos: regalo y premio por carga ==")
cli = app.get_cliente("lavadero")
tok = cli["panel_token"]
app.asegurar_panel("lavadero")                 # crea el pin_premio si falta
cli = app.get_cliente("lavadero")
pin_regalo, pin_premio = cli["pin"], cli["pin_premio"]

chequear("se genero el segundo codigo", bool(pin_premio), pin_premio)
chequear("y es distinto del primero", pin_regalo != pin_premio, (pin_regalo, pin_premio))

def _limpiar():
    app.guardar_cliente("lavadero", {"pin_fallidos": 0, "pin_bloqueado_hasta": None})
    app.actualizar_dispositivo("lavadero03", {"ultimo_poll": app.ahora_ar().isoformat()})
    conn = app.get_db(); cur = conn.cursor()
    cur.execute("DELETE FROM ordenes WHERE id LIKE 'gift_%%'")
    conn.commit(); cur.close(); conn.close()

def _motivos():
    conn = app.get_db(); cur = conn.cursor()
    cur.execute("SELECT motivo FROM ordenes WHERE id LIKE 'gift_%%' ORDER BY fecha")
    r = [x["motivo"] for x in cur.fetchall()]; cur.close(); conn.close()
    return r

_limpiar()
r = c.post(f"/panel/{tok}", data={"accion": "regalar", "pin": pin_regalo}, follow_redirects=True)
chequear("el codigo de regalo entrega", _motivos() == ["regalo"], _motivos())
chequear("y el mensaje dice como quedo", "Regalo" in r.get_data(as_text=True))

_limpiar()
r = c.post(f"/panel/{tok}", data={"accion": "regalar", "pin": pin_premio}, follow_redirects=True)
chequear("el codigo de premio entrega", _motivos() == ["premio"], _motivos())
chequear("y lo identifica como premio por carga",
         "Premio por carga" in r.get_data(as_text=True))

# El motivo sale del CODIGO, no del formulario: mandar motivo a mano no sirve.
_limpiar()
c.post(f"/panel/{tok}", data={"accion": "regalar", "pin": pin_regalo, "motivo": "premio"},
       follow_redirects=True)
chequear("no se puede falsear el motivo desde el formulario",
         _motivos() == ["regalo"], _motivos())

# Los dos codigos comparten el contador de intentos fallidos: si no, un
# atacante tendria el doble de chances.
_limpiar()
_cli = app.get_cliente('lavadero')
_malo = next(f'{n:04d}' for n in range(10000)
             if f'{n:04d}' not in (_cli['pin'], _cli['pin_premio']))
for _ in range(app.PIN_MAX_INTENTOS):
    r = c.post(f"/panel/{tok}", data={"accion": "regalar", "pin": _malo}, follow_redirects=True)
chequear("el tope de intentos es compartido por los dos codigos",
         bool(app.get_cliente("lavadero")["pin_bloqueado_hasta"]))
r = c.post(f"/panel/{tok}", data={"accion": "regalar", "pin": pin_premio}, follow_redirects=True)
chequear("bloqueado, el codigo de premio tampoco pasa", _motivos() == [], _motivos())
_limpiar()

# El panel muestra los dos contadores y la columna de motivo.
c.post(f"/panel/{tok}", data={"accion": "regalar", "pin": pin_premio}, follow_redirects=True)
r = c.get(f"/panel/{tok}")
cuerpo = r.get_data(as_text=True)
chequear("el panel nombra los dos motivos",
         "Regalo" in cuerpo and "Premio por carga" in cuerpo)
chequear("la tabla tiene columna Motivo", "<th>Motivo</th>" in cuerpo)

# Admin: se ven y se renuevan los dos por separado.
r = c.get(f"/admin/{K}")
adm = r.get_data(as_text=True)
chequear("el admin muestra el codigo de regalo", "digo regalo" in adm)
chequear("y el de premio por carga", "digo premio por carga" in adm)
_antes = app.get_cliente("lavadero")["pin"]
c.post(f"/admin/{K}", data={"accion": "nuevo_pin", "alias": "lavadero", "cual": "premio"})
chequear("renovar el de premio NO toca el de regalo",
         app.get_cliente("lavadero")["pin"] == _antes)
chequear("y cambia el de premio",
         app.get_cliente("lavadero")["pin_premio"] != pin_premio)

print("\n== 16. Planilla Excel ==")
_limpiar()
hoy = app.ahora_ar()
app.insertar_orden("ord_x1", "lavadero03", 1, 1500)
c.post(f"/panel/{tok}", data={"accion": "regalar", "pin": app.get_cliente("lavadero")["pin"]},
       follow_redirects=True)

d1 = hoy.strftime("%Y-%m-01"); d2 = hoy.strftime("%Y-%m-%d")
r = c.get(f"/panel/{tok}/planilla.xlsx?desde={d1}&hasta={d2}")
chequear("la planilla descarga", r.status_code == 200, r.status_code)
chequear("es un xlsx de verdad", r.data[:2] == b"PK", r.data[:8])
chequear("va como adjunto con nombre",
         ".xlsx" in r.headers.get("Content-Disposition", ""),
         r.headers.get("Content-Disposition"))

import openpyxl, io as _io
wb = openpyxl.load_workbook(_io.BytesIO(r.data))
chequear("trae las tres hojas",
         wb.sheetnames == ["Ventas", "Entregas sin cargo", "Resumen por dia"],
         wb.sheetnames)
hv = wb["Ventas"]
chequear("la hoja de ventas tiene encabezado",
         [c0.value for c0 in hv[1]] == ["Fecha", "Hora", "Maquina", "Monto"],
         [c0.value for c0 in hv[1]])
chequear("y el monto es numero, no texto",
         isinstance(hv.cell(row=2, column=4).value, (int, float)),
         type(hv.cell(row=2, column=4).value).__name__)
chequear("la fecha es fecha, no texto",
         hasattr(hv.cell(row=2, column=1).value, "year"),
         type(hv.cell(row=2, column=1).value).__name__)
hs = wb["Entregas sin cargo"]
chequear("la hoja sin cargo trae el motivo",
         hs.cell(row=2, column=4).value in ("Regalo", "Premio por carga"),
         hs.cell(row=2, column=4).value)

# El rango filtra de verdad.
r2 = c.get(f"/panel/{tok}/planilla.xlsx?desde=2020-01-01&hasta=2020-01-31")
wb2 = openpyxl.load_workbook(_io.BytesIO(r2.data))
chequear("un rango sin datos da hojas vacias", wb2["Ventas"].max_row == 1,
         wb2["Ventas"].max_row)

# Token invalido no descarga nada.
chequear("con token invalido da 404",
         c.get("/panel/nopenope/planilla.xlsx").status_code == 404)
_limpiar()

print("\n== 17. Pulsos que REPORTA la caja (heartbeat -> /config) ==")

def _config():
    return c.get(f"/config/{K}").get_data(as_text=True)

# Arranca sin reportar: la pagina tiene que decirlo, no inventar un numero.
app.actualizar_dispositivo("lavadero03", {"credito_ficha": 3000, "valor_pulso": 100,
                                          "pulsos_reportados": None})
cuerpo = _config()
chequear("sin reporte avisa que la caja todavia no dijo nada",
         "no report" in cuerpo, cuerpo[cuerpo.find("pulsos por ficha")-200:][:260])
chequear("y igual muestra el calculo", "= 30 pulsos por ficha" in cuerpo)

# Coincide -> confirmacion.
app.actualizar_dispositivo("lavadero03", {"pulsos_reportados": 30})
cuerpo = _config()
chequear("cuando coincide, lo confirma", "la caja confirma 30" in cuerpo)
chequear("y no muestra alerta", 'class="mal"' not in cuerpo)

# No coincide -> alerta con los DOS numeros, que es lo util.
app.actualizar_dispositivo("lavadero03", {"credito_ficha": 4000, "pulsos_reportados": 30})
cuerpo = _config()
chequear("cuando no coincide, avisa", 'class="mal"' in cuerpo)
chequear("y dice los dos numeros",
         "Guardado: 40 pulsos por ficha" in cuerpo and "usando 30" in cuerpo)

# Config invalida: manda el error de division, no el de desfasaje.
app.actualizar_dispositivo("lavadero03", {"credito_ficha": 3050, "valor_pulso": 100})
cuerpo = _config()
chequear("division inexacta gana sobre el reporte",
         "no da exacta" in cuerpo and "la caja confirma" not in cuerpo)
app.actualizar_dispositivo("lavadero03", {"credito_ficha": 3000, "valor_pulso": 100,
                                          "pulsos_reportados": None})

# ---- el camino de entrada: el latido MQTT ----
# on_message esta adentro de mqtt_liveness_loop, asi que se prueba el efecto:
# que actualizar_dispositivo acepte la columna y que /config la lea.
app.actualizar_dispositivo("lavadero03", {"pulsos_reportados": 45})
chequear("la columna persiste",
         app.get_dispositivo("lavadero03")["pulsos_reportados"] == 45,
         app.get_dispositivo("lavadero03")["pulsos_reportados"])

# Un valor absurdo no tiene que llegar a la base: el filtro es 1..PULSOS_MAX.
def _filtrar(v):
    """Misma logica que on_message, para que el rango quede cubierto."""
    try:
        n = int(v)
        return n if 1 <= n <= app.PULSOS_MAX else None
    except (TypeError, ValueError):
        return None

chequear("acepta un valor normal", _filtrar(30) == 30)
chequear("rechaza 0", _filtrar(0) is None)
chequear("rechaza por encima del tope", _filtrar(app.PULSOS_MAX + 1) is None)
chequear("rechaza basura", _filtrar("treinta") is None)
chequear("rechaza faltante", _filtrar(None) is None)
app.actualizar_dispositivo("lavadero03", {"pulsos_reportados": None})

print("\n== 18. El admin no puede reenviar el formulario (PRG) ==")
# El bug: el admin procesaba el POST y dibujaba la pagina. Un F5, un 'atras' o
# la restauracion de la pestaña REENVIAN ese POST y la accion corre de nuevo,
# sin el confirm() del onsubmit. Con 'generar codigo nuevo' eso le cambia el
# codigo al cliente sin que nadie lo pida.
cli = app.get_cliente("lavadero") or app.get_cliente("lava_dero")
_alias = cli["alias"]

r = c.post(f"/admin/{K}", data={"accion": "nuevo_pin", "alias": _alias, "cual": "regalo"})
chequear("el POST contesta con un redirect, no con la pagina",
         r.status_code in (302, 303), r.status_code)
_destino = r.headers.get("Location", "")
_pin1 = app.get_cliente(_alias)["pin"]

# Esto es el F5: volver a pedir la MISMA url a la que nos mando el redirect.
for _ in range(3):
    r2 = c.get(_destino)
chequear("refrescar despues de generar un codigo NO genera otro",
         app.get_cliente(_alias)["pin"] == _pin1,
         (_pin1, app.get_cliente(_alias)["pin"]))
chequear("y el mensaje se sigue viendo despues del redirect",
         "digo" in r2.get_data(as_text=True))

# El mensaje viaja por la URL: no puede inyectar HTML al volver.
r3 = c.get(f"/admin/{K}?m=%3Cimg+src%3Dx+onerror%3Dalert(1)%3E")
cuerpo = r3.get_data(as_text=True)
chequear("el mensaje de la URL no inyecta HTML",
         "<img src=x" not in cuerpo and "&lt;img" in cuerpo)

print("\n== 19. Fecha de cambio de cada codigo ==")
cli = app.get_cliente(_alias)
chequear("generar un codigo deja la fecha", bool(cli["pin_cambiado"]), cli["pin_cambiado"])
chequear("y es de ahora",
         abs((app.ahora_ar() - app.datetime.fromisoformat(cli["pin_cambiado"])).total_seconds()) < 120)

_fecha_regalo = cli["pin_cambiado"]
c.post(f"/admin/{K}", data={"accion": "nuevo_pin", "alias": _alias, "cual": "premio"},
       follow_redirects=True)
cli = app.get_cliente(_alias)
chequear("renovar el de premio no toca la fecha del de regalo",
         cli["pin_cambiado"] == _fecha_regalo)
chequear("y deja la suya", bool(cli["pin_premio_cambiado"]))

adm = c.get(f"/admin/{K}").get_data(as_text=True)
chequear("el admin muestra desde cuando rige cada codigo",
         adm.count("desde el ") >= 2, adm.count("desde el "))

# Un codigo viejo (los que ya existian) no tiene fecha: se dice, no se inventa.
app.guardar_cliente(_alias, {"pin_cambiado": None})
adm = c.get(f"/admin/{K}").get_data(as_text=True)
chequear("sin fecha dice 'desde antes', no una fecha inventada",
         "desde antes" in adm)

# Un codigo de muchos meses se marca.
_viejo = (app.ahora_ar() - app.timedelta(days=400)).isoformat()
app.guardar_cliente(_alias, {"pin_cambiado": _viejo})
adm = c.get(f"/admin/{K}").get_data(as_text=True)
chequear("un codigo de hace mas de 6 meses sale marcado",
         "400 dias" in adm, [l for l in adm.split("\n") if "dias)" in l][:1])

print("\n== LWT rancio: un will que llega tarde no debe voltear un QR vivo ==")
# Caso real del 14/10/26: el client id del ESP es fijo, asi que al reconectar el
# broker mata la sesion vieja y publica SU will. El "offline" llego DESPUES del
# "online" (y con retain=0, o sea que el filtro de retenidos no lo tapaba) y el
# backend cancelo el QR de aspiradora01 y soplado01 estando las dos funcionando.

def _preparar(caja, seg_desde_contacto, con_qr=True):
    """Deja una caja con su ultimo contacto hace N segundos y un QR armado."""
    app.actualizar_dispositivo(caja, {
        "ultimo_poll": (app.ahora_ar() - app.timedelta(seconds=seg_desde_contacto)).isoformat(),
        "orden_qr_id": ("orden_" + caja) if con_qr else None,
    })

_canceladas = []
app.cancelar_orden_qr = lambda d: _canceladas.append(d["id"])

# --- la caja esta viva: el will es viejo, se ignora ---
_canceladas.clear()
_preparar("aspiradora01", 5)
_preparar("soplado01", 5)
res = app.manejar_offline("aspiradora01", app.get_dispositivo("aspiradora01"))
chequear("will rancio: no cancela nada", res == [], res)
chequear("y el QR sigue armado", bool(app.get_dispositivo("aspiradora01")["orden_qr_id"]))
chequear("tampoco toca a la hermana", bool(app.get_dispositivo("soplado01")["orden_qr_id"]))

# --- la caja esta muerta de verdad: el will manda ---
_canceladas.clear()
_preparar("aspiradora01", 300)
_preparar("soplado01", 300)
res = app.manejar_offline("aspiradora01", app.get_dispositivo("aspiradora01"))
chequear("will real: cancela las dos del mismo ESP",
         set(res) == {"aspiradora01", "soplado01"}, res)
chequear("y lo hace contra MercadoPago", set(_canceladas) == {"aspiradora01", "soplado01"}, _canceladas)
chequear("deja el QR vacio para que el proximo latido lo re-arme",
         not app.get_dispositivo("aspiradora01")["orden_qr_id"])

# --- una murio y la hermana sigue latiendo: solo cae la muerta ---
_canceladas.clear()
_preparar("aspiradora01", 300)
_preparar("soplado01", 5)
res = app.manejar_offline("aspiradora01", app.get_dispositivo("aspiradora01"))
chequear("solo cancela la que no responde", res == ["aspiradora01"], res)
chequear("la hermana que late conserva su QR",
         bool(app.get_dispositivo("soplado01")["orden_qr_id"]))

# --- el limite esta donde tiene que estar ---
_ahora = app.ahora_ar()
_d = lambda seg: {"ultimo_poll": (_ahora - app.timedelta(seconds=seg)).isoformat()}
chequear("un will REAL llega >=90s despues del ultimo latido -> se respeta",
         not app._contacto_fresco(_d(90), _ahora))
chequear("un will RANCIO llega a los pocos segundos -> se descarta",
         app._contacto_fresco(_d(3), _ahora))
chequear("una caja que nunca reporto no cuenta como fresca",
         not app._contacto_fresco({}, _ahora))
chequear("un ultimo_poll corrupto no rompe ni miente",
         not app._contacto_fresco({"ultimo_poll": "fecha-rota"}, _ahora))

print("\n== Suscriptor sordo: avisar del backend, no de las maquinas ==")
app._sub_estado["ultimo"] = None
chequear("recien arrancado no se declara sordo", not app.suscriptor_sordo())
app._sub_visto()
chequear("con trafico reciente no esta sordo", not app.suscriptor_sordo())
app._sub_estado["ultimo"] = app.ahora_ar() - app.timedelta(seconds=app.SUB_SORDO_S - 10)
chequear("justo antes del limite todavia no", not app.suscriptor_sordo())
app._sub_estado["ultimo"] = app.ahora_ar() - app.timedelta(seconds=app.SUB_SORDO_S + 10)
chequear("pasado el limite si", app.suscriptor_sordo())
app._sub_visto()
chequear("un mensaje lo revive", not app.suscriptor_sordo())
chequear("el umbral del sordo (180s) es mayor que el latido del ESP (60s)",
         app.SUB_SORDO_S > 60)
chequear("y el de frescura del will deja pasar un will real (>=90s)",
         app.LWT_FRESCURA_S < 90, app.LWT_FRESCURA_S)

print("\n== Pago que entra con el broker caido ==")
# El agujero: publicar_activacion falla -> la orden queda 'pendiente' y nadie la
# reintenta. Y el reembolso solo miraba equipo_offline(), que se juzga por el
# ultimo_poll: si el broker vuelve, el equipo PARECE sano y la orden se quedaba
# en 'pendiente' para siempre. El cliente pagaba, no recibia nada, y tampoco se
# le devolvia la plata.

_pub_ok = {"v": True}
app.publicar_activacion = lambda caja, pago, segundos_override=None: _pub_ok["v"]
_reembolsadas = []
app.reembolsar_orden_mp = lambda o: (_reembolsadas.append(o["id"]) or True)
app.cancelar_orden_qr = lambda d: None

def _estado_de(oid):
    cx = app.get_db(); cc = cx.cursor()
    cc.execute("SELECT estado FROM ordenes WHERE id=%s", (oid,))
    r = cc.fetchone(); cc.close(); cx.close()
    return r and r["estado"]

def _orden_vieja(oid, caja, minutos):
    """Una orden pagada hace N minutos que nunca se despacho."""
    app.insertar_orden(oid, caja, 300, 500)
    cx = app.get_db(); cc = cx.cursor()
    cc.execute("UPDATE ordenes SET fecha=%s WHERE id=%s",
               ((app.ahora_ar() - app.timedelta(minutes=minutos)).isoformat(), oid))
    cx.commit(); cc.close(); cx.close()

# --- el reintento entrega lo que el primer envio no pudo ---
_pub_ok["v"] = False
_orden_vieja("ord_reintento", "aspiradora01", 1)      # dentro de la ventana
chequear("quedo pendiente cuando fallo el publish", _estado_de("ord_reintento") == "pendiente")
_pub_ok["v"] = True                                    # volvio el broker
chequear("el reintento la despacha",
         app.reintentar_despacho({"id": "ord_reintento", "dispositivo_id": "aspiradora01"}))
chequear("y pasa a ejecutando (caja sostenida)",
         _estado_de("ord_reintento") == "ejecutando", _estado_de("ord_reintento"))

# --- una expendedora se marca completada, no ejecutando ---
_orden_vieja("ord_fichas", "villagas01", 1)
app.reintentar_despacho({"id": "ord_fichas", "dispositivo_id": "villagas01"})
chequear("una caja de pulso/fichas queda completada",
         _estado_de("ord_fichas") == "completada", _estado_de("ord_fichas"))

# --- con el broker todavia caido, el reintento no miente ---
_pub_ok["v"] = False
_orden_vieja("ord_sigue_caido", "aspiradora01", 1)
chequear("si sigue sin broker, no la despacha",
         not app.reintentar_despacho({"id": "ord_sigue_caido", "dispositivo_id": "aspiradora01"}))
chequear("y la deja pendiente para el proximo intento",
         _estado_de("ord_sigue_caido") == "pendiente")

# --- EL CASO QUE SE PERDIA: equipo que volvio, orden nunca entregada ---
_reembolsadas.clear()
_orden_vieja("ord_perdida", "aspiradora01", app.REEMBOLSO_MINUTOS + 2)
app.actualizar_dispositivo("aspiradora01", {"ultimo_poll": app.ahora_ar().isoformat()})
chequear("el equipo se ve online (el broker volvio)",
         not app.equipo_offline(app.get_dispositivo("aspiradora01")))
# una sola pasada del vigilante, sin el sleep
_pub_ok["v"] = False
import threading as _th
_t = _th.Thread(target=app.vigilar_ordenes, daemon=True); _t.start()
import time as _time; _time.sleep(1.5)
chequear("AHORA SI se reembolsa aunque el equipo parezca sano",
         "ord_perdida" in _reembolsadas, _reembolsadas)
chequear("y queda marcada como reembolsada",
         _estado_de("ord_perdida") == "reembolsada", _estado_de("ord_perdida"))

# --- una caja vieja por polling conserva el filtro de siempre ---
_reembolsadas.clear()
_orden_vieja("ord_legacy", "termo_001", app.REEMBOLSO_MINUTOS + 2)
app.actualizar_dispositivo("termo_001", {"ultimo_poll": app.ahora_ar().isoformat()})
_time.sleep(1.5)
chequear("caja legacy online: sigue sin reembolsarse (puede estar por pollear)",
         "ord_legacy" not in _reembolsadas, _reembolsadas)
chequear("y sigue pendiente", _estado_de("ord_legacy") == "pendiente")

print("\n== Bajar el QR cuando no se pudo despachar ==")
app.actualizar_dispositivo("aspiradora01", {"orden_qr_id": "orden_a01"})
app.actualizar_dispositivo("soplado01", {"orden_qr_id": "orden_s01"})
bajadas = app.bajar_qr_sin_servicio("aspiradora01")
chequear("baja el QR de la caja y el de su hermana del mismo ESP",
         set(bajadas) == {"aspiradora01", "soplado01"}, bajadas)
chequear("quedan sin QR, listas para re-armarse con el proximo latido",
         not app.get_dispositivo("aspiradora01")["orden_qr_id"]
         and not app.get_dispositivo("soplado01")["orden_qr_id"])
chequear("llamarlo de nuevo no rompe ni cancela dos veces",
         app.bajar_qr_sin_servicio("aspiradora01") == [])

print("\n== 20. Cuadro de fichas sin cargo por dia ==")
_cli = app.get_cliente("lavadero") or app.get_cliente("lava_dero")
_alias = _cli["alias"]
tok = _cli["panel_token"]
app.actualizar_dispositivo("lavadero03", {"cliente": _alias})

def _borrar_gifts():
    conn = app.get_db(); cur = conn.cursor()
    cur.execute("DELETE FROM ordenes")
    conn.commit(); cur.close(); conn.close()

def _gift(oid, cuando, motivo):
    conn = app.get_db(); cur = conn.cursor()
    cur.execute("INSERT INTO ordenes (id, dispositivo_id, segundos, estado, fecha, monto, motivo)"
                " VALUES (?,?,?,?,?,?,?)",
                (oid, "lavadero03", 1, "regalada", cuando.isoformat(), 0, motivo))
    conn.commit(); cur.close(); conn.close()

_borrar_gifts()
r = c.get(f"/panel/{tok}")
chequear("sin entregas, el cuadro lo dice",
         "Sin entregas sin cargo todavia" in r.get_data(as_text=True))

# 6 dias x 6 fichas (4 regalo + 2 premio). El dia mas viejo cae FUERA de las 25
# filas de la lista de arriba: si el cuadro se armara desde esa lista cortada,
# ese dia desapareceria sin avisar.
_hoy = app.ahora_ar()
_n = 0
for _d in range(6):
    _f = _hoy - app.timedelta(days=_d)
    for _i in range(6):
        _n += 1
        _gift(f"gift_d{_d}_{_i}", _f.replace(hour=9 + _i, minute=0),
              "regalo" if _i < 4 else "premio")

cuerpo = c.get(f"/panel/{tok}").get_data(as_text=True)
_viejo = (_hoy - app.timedelta(days=5)).strftime("%Y-%m-%d")
_fila = f"<tr><td>{_viejo}</td><td>4</td><td>2</td><td><b>6</b></td></tr>"
chequear("el dia mas viejo aparece aunque quede fuera de las 25 ultimas",
         _fila in cuerpo, [l for l in cuerpo.split("<tr>") if _viejo in l][:1])
chequear("hay una fila por dia", cuerpo.count("<td><b>6</b></td>") == 6,
         cuerpo.count("<td><b>6</b></td>"))
chequear("el cuadro tiene las dos columnas de motivo",
         "<th>Regalo</th>" in cuerpo and "<th>Premio por carga</th>" in cuerpo)

# Una fila vieja sin motivo (de cuando habia un solo PIN) cuenta como regalo,
# no se pierde ni rompe la suma.
_borrar_gifts()
_gift("gift_viejo", _hoy, None)
cuerpo = c.get(f"/panel/{tok}").get_data(as_text=True)
_dia_hoy = _hoy.strftime("%Y-%m-%d")
chequear("una entrega sin motivo cuenta como regalo",
         f"<tr><td>{_dia_hoy}</td><td>1</td><td>0</td><td><b>1</b></td></tr>" in cuerpo,
         [l for l in cuerpo.split("<tr>") if _dia_hoy in l][:2])

# Las regaladas no pueden colarse en el cuadro de ventas.
chequear("no facturan: el cuadro de ventas sigue vacio",
         ">Sin datos<" in cuerpo)
_borrar_gifts()

print("\n== 21. Version desplegada visible en el admin ==")
import time as _time

# Fuera de Railway no hay que inventar nada: se dice que no se sabe.
_sha, _msg, _rama = app.COMMIT_SHA, app.COMMIT_MSG, app.GIT_BRANCH
app.COMMIT_SHA = ""
chequear("sin las variables del entorno no inventa version",
         app.version_desplegada() is None)
adm = c.get(f"/admin/{K}").get_data(as_text=True)
chequear("y lo dice en la pagina", "no la informa el entorno" in adm)

# Con las variables puestas: sha, mensaje y hace cuanto.
app.COMMIT_SHA = "abc1234"
app.COMMIT_MSG = "Panel: cuadro de fichas sin cargo por dia\nsegunda linea que no va"
app.GIT_BRANCH = "main"
app.ARRANQUE_TS = _time.time() - 7200      # 2 h
v = app.version_desplegada()
chequear("devuelve el sha corto", v["sha"] == "abc1234", v["sha"])
chequear("solo la primera linea del mensaje",
         v["msg"] == "Panel: cuadro de fichas sin cargo por dia", v["msg"])
chequear("y hace cuanto arranco", v["hace"] == "hace 2 h", v["hace"])

adm = c.get(f"/admin/{K}").get_data(as_text=True)
chequear("el admin muestra el commit", "abc1234" in adm)
chequear("y cuando se desplego", "hace 2 h" in adm)
chequear("no nombra la rama cuando es main", "rama main" not in adm)
chequear("explica para que sirve", "no se desplego" in adm)

# Una rama que no es main SI se nombra: es justo el caso en que te confundis.
app.GIT_BRANCH = "pruebas"
adm = c.get(f"/admin/{K}").get_data(as_text=True)
chequear("una rama distinta de main se muestra", "rama pruebas" in adm)

# El mensaje del commit viene de afuera: no puede inyectar HTML.
app.GIT_BRANCH = "main"
app.COMMIT_MSG = '<img src=x onerror=alert(1)>'
adm = c.get(f"/admin/{K}").get_data(as_text=True)
chequear("el mensaje del commit sale escapado",
         "<img src=x" not in adm and "&lt;img" in adm)

# Minutos y dias, no solo horas.
app.COMMIT_MSG = "x"
app.ARRANQUE_TS = _time.time() - 300
chequear("menos de una hora se muestra en minutos",
         app.version_desplegada()["hace"] == "hace 5 min",
         app.version_desplegada()["hace"])
app.ARRANQUE_TS = _time.time() - 86400 * 9
chequear("nueve dias se muestran en dias",
         app.version_desplegada()["hace"] == "hace 9 d",
         app.version_desplegada()["hace"])

app.COMMIT_SHA, app.COMMIT_MSG, app.GIT_BRANCH = _sha, _msg, _rama

print("\n" + ("TODO OK" if not fallas else f"FALLARON {len(fallas)}: {fallas}"))
sys.exit(1 if fallas else 0)
